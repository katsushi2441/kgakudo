#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""こども家庭庁「保育所等関連状況取りまとめ」の Excel を、住所から引ける形にする。

**学童（放課後児童クラブ）との違い**
- 学童は公表がPDFだけで、市区町村名で数字が出るのは待機50人以上の88市町村だけだった。
- 保育所は **Excel で 1,741市区町村ぶん**（定員・申込者）が出ていて、待機児童がいる市区町村も
  資料6-1/6-2 に一覧がある。**国が機械可読で出している数少ない例**。
  なので当社の仕事は「PDFを開く」ことではなく、**住所から引ける形にする**ことだけ。

取り込むもの（令和7年4月1日現在）:
  定員・申込者の状況.xlsx  … 1,741市区町村の 定員（保育所・認定こども園・地域型など）と申込者数
  資料1〜6.xlsx の 資料6-1/6-2 … 待機児童がいる市区町村（211件・今年と前年）

  /usr/bin/python3 scripts/parse_hoiku.py 定員申込.xlsx 資料1-6.xlsx --year 2025

**推定はしない。** 待機児童の一覧に無い市区町村は「待機なし（国の一覧に載っていない）」として
waiting を 0 ではなく None で持ち、画面側で「一覧に載っていません」と書けるようにする。
"""
import argparse
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE = 'こども家庭庁「保育所等関連状況取りまとめ」'
SOURCE_URL = 'https://www.cfa.go.jp/policies/hoiku/torimatome/'


def _int(v):
    try:
        return int(str(v).replace(',', '').strip())
    except (TypeError, ValueError):
        return None


def parse_capacity(path: str) -> dict:
    """定員・申込者の状況。シート2枚を都道府県＋市区町村で突き合わせる。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for sheet, keys in (('定員の状況', ('capacity_hoikusho', 'capacity_kodomoen')),
                        ('申込者の状況', ('applicants', 'using_hoikusho'))):
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            pref, city = row[2], row[3]
            if not pref or not city or not str(pref).endswith(('都', '道', '府', '県')):
                continue
            a, b = _int(row[4]), _int(row[5])
            if a is None:
                continue
            rec = out.setdefault((str(pref), str(city)), dict(pref=str(pref), city=str(city)))
            rec[keys[0]] = a
            rec[keys[1]] = b
    return out


def parse_waiting(path: str) -> list:
    """資料6-1/6-2。左右2段組みなので、両方の列の組を拾う。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, seen = [], set()
    for name in wb.sheetnames:
        if '６－' not in name and '6－' not in name and '6-' not in name:
            continue
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            cells = list(row) + [None] * 16
            for base in (2, 8):   # 左の段 / 右の段
                pref, city, now, prev = cells[base], cells[base + 1], cells[base + 2], cells[base + 3]
                if not pref or not city:
                    continue
                if not str(pref).endswith(('都', '道', '府', '県')):
                    continue
                n, p = _int(now), _int(prev)
                if n is None:
                    continue
                key = (str(pref), str(city))
                if key in seen:
                    continue
                seen.add(key)
                rows.append(dict(pref=str(pref), city=str(city), waiting=n, waiting_prev=p))
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('capacity_xlsx')
    ap.add_argument('shiryo_xlsx')
    ap.add_argument('--year', type=int, required=True)
    a = ap.parse_args()

    cap = parse_capacity(a.capacity_xlsx)
    wait = parse_waiting(a.shiryo_xlsx)
    by = {(w['pref'], w['city']): w for w in wait}
    for key, rec in cap.items():
        w = by.get(key)
        rec['waiting'] = w['waiting'] if w else None
        rec['waiting_prev'] = w['waiting_prev'] if w else None

    municipalities = sorted(cap.values(), key=lambda r: (r['pref'], r['city']))
    listed = [r for r in municipalities if r['waiting'] is not None]
    out = dict(year=a.year, as_of=f'{a.year}-04-01', source=SOURCE, source_url=SOURCE_URL,
               note='定員・申込者は全1,741市区町村。待機児童数は「待機児童がいる市区町村」の一覧にある分だけで、'
                    '一覧に無い市区町村は waiting を null にしてある（0人と断定しない）。',
               municipalities=municipalities)
    p = os.path.join(ROOT, 'data', f'hoiku_{a.year}.json')
    with open(p, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f'市区町村 {len(municipalities)} 件 / 待機児童の記載あり {len(listed)} 件')
    print('待機児童の合計:', sum(r["waiting"] for r in listed))
    print('→', p)
    if len(municipalities) < 1700:
        print('！ 市区町村が少なすぎる。列の位置が変わった可能性がある', file=sys.stderr)


if __name__ == '__main__':
    main()
